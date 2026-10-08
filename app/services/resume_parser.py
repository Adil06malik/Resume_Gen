import fitz
import docx
from pathlib import Path
from typing import Optional
from app.prompts import RESUME_PARSER_PROMPT
from app.ai.factory import get_ai_provider
from app.schemas.resume import StructuredResume
import json


class ResumeParserService:
    def __init__(self):
        try:
            self.ai_provider = get_ai_provider()
        except ValueError:
            self.ai_provider = None

    async def extract_text(self, file_path: str) -> str:
        path = Path(file_path)
        suffix = path.suffix.lower()

        if suffix == ".pdf":
            return self._extract_pdf(file_path)
        elif suffix == ".docx":
            return self._extract_docx(file_path)
        elif suffix == ".txt":
            return self._extract_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {suffix}")

    def _extract_pdf(self, file_path: str) -> str:
        text = ""
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text()
        return text

    def _extract_docx(self, file_path: str) -> str:
        doc = docx.Document(file_path)
        text = []
        for para in doc.paragraphs:
            text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text.append(cell.text)
        return "\n".join(text)

    def _extract_txt(self, file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()

    async def parse_resume(self, text: str) -> StructuredResume:
        if not self.ai_provider:
            raise ValueError("AI provider not configured. Set AI_API_KEY in .env")
        prompt = RESUME_PARSER_PROMPT.replace("{{resume_text}}", text)

        max_retries = 2
        for attempt in range(max_retries):
            try:
                result = await self.ai_provider.generate_json(prompt)
                return StructuredResume(**result)
            except (json.JSONDecodeError, ValueError, TypeError) as e:
                if attempt == max_retries - 1:
                    raise ValueError(f"Failed to parse resume after {max_retries} attempts: {e}")
                continue

        raise ValueError("Failed to parse resume")

    async def close(self):
        if self.ai_provider:
            await self.ai_provider.close()