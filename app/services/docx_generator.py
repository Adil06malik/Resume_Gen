import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.config import settings
from app.schemas.resume import StructuredResume


class DOCXGeneratorService:
    def __init__(self):
        os.makedirs(settings.GENERATED_DIR, exist_ok=True)

    def generate_resume_docx(
        self,
        resume: StructuredResume,
        company: str,
        job_title: str,
    ) -> str:
        safe_company = "".join(c if c.isalnum() else "_" for c in company)
        safe_title = "".join(c if c.isalnum() else "_" for c in job_title)
        filename = f"Resume_{safe_company}_{safe_title}.docx"
        filepath = os.path.join(settings.GENERATED_DIR, filename)

        doc = Document()

        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        pf = style.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)

        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        c = resume.candidate
        if c.name:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(c.name)
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

        contact_parts = []
        if c.email:
            contact_parts.append(c.email)
        if c.phone:
            contact_parts.append(c.phone)
        if c.location:
            contact_parts.append(c.location)
        if contact_parts:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(" | ".join(contact_parts))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        if c.links:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(" | ".join(c.links))
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        def add_section_heading(text):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
            border_p = doc.add_paragraph()
            border_p.paragraph_format.space_before = Pt(0)
            border_p.paragraph_format.space_after = Pt(4)

        if resume.summary:
            add_section_heading("PROFESSIONAL SUMMARY")
            p = doc.add_paragraph(resume.summary)

        if resume.skills:
            add_section_heading("SKILLS")
            p = doc.add_paragraph(", ".join(s.name for s in resume.skills))

        if resume.experience:
            add_section_heading("EXPERIENCE")
            for exp in resume.experience:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(exp.title)
                run.bold = True
                run = p.add_run(f" - {exp.company}")
                if exp.location:
                    run = p.add_run(f", {exp.location}")

                date_range = exp.start_date
                if exp.end_date:
                    date_range += f" - {exp.end_date}"
                p = doc.add_paragraph()
                run = p.add_run(date_range)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

                for desc in exp.description:
                    p = doc.add_paragraph(desc, style="List Bullet")

                if exp.technologies:
                    p = doc.add_paragraph()
                    run = p.add_run("Technologies: ")
                    run.bold = True
                    run.font.size = Pt(9)
                    run = p.add_run(", ".join(exp.technologies))
                    run.font.size = Pt(9)

        if resume.projects:
            add_section_heading("PROJECTS")
            for proj in resume.projects:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                run = p.add_run(proj.name)
                run.bold = True
                if proj.url:
                    run = p.add_run(f" - {proj.url}")

                dates = ""
                if proj.start_date:
                    dates = proj.start_date
                    if proj.end_date:
                        dates += f" - {proj.end_date}"
                if dates:
                    p = doc.add_paragraph()
                    run = p.add_run(dates)
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

                if proj.description:
                    doc.add_paragraph(proj.description)

                if proj.technologies:
                    p = doc.add_paragraph()
                    run = p.add_run("Technologies: ")
                    run.bold = True
                    run.font.size = Pt(9)
                    run = p.add_run(", ".join(proj.technologies))
                    run.font.size = Pt(9)

        if resume.education:
            add_section_heading("EDUCATION")
            for edu in resume.education:
                p = doc.add_paragraph()
                run = p.add_run(edu.degree)
                run.bold = True
                run = p.add_run(f" - {edu.institution}")
                if edu.location:
                    run = p.add_run(f", {edu.location}")

                meta_parts = []
                if edu.graduation_date:
                    meta_parts.append(edu.graduation_date)
                if edu.gpa:
                    meta_parts.append(f"GPA: {edu.gpa}")
                if meta_parts:
                    p = doc.add_paragraph()
                    run = p.add_run(" | ".join(meta_parts))
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        if resume.certifications:
            add_section_heading("CERTIFICATIONS")
            for cert in resume.certifications:
                p = doc.add_paragraph()
                run = p.add_run(cert.name)
                run.bold = True
                run = p.add_run(f" - {cert.issuer}")
                if cert.date:
                    run = p.add_run(f" ({cert.date})")

        doc.save(filepath)
        return filepath
