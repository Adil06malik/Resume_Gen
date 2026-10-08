import os
import tempfile
from unittest.mock import patch, MagicMock

import pytest
import fitz
import docx

from app.services.resume_parser import ResumeParserService


class FakeProvider:
    async def generate(self, prompt, system_prompt=None):
        return ""

    async def generate_json(self, prompt, system_prompt=None):
        return {}

    async def close(self):
        pass


@pytest.fixture
def sample_pdf():
    pdf_path = tempfile.mktemp(suffix=".pdf")
    doc = fitz.open()
    page = doc.new_page()
    text = (
        "John Doe\n"
        "john@example.com | 555-0100 | New York, NY\n"
        "\n"
        "PROFESSIONAL SUMMARY\n"
        "Experienced software engineer with 5+ years in web development.\n"
        "\n"
        "SKILLS\n"
        "Python, FastAPI, PostgreSQL, Docker\n"
        "\n"
        "EXPERIENCE\n"
        "Senior Developer | Tech Corp | Jan 2020 - Present\n"
        "- Led development of microservices architecture\n"
        "- Improved API response time by 40%\n"
        "\n"
        "EDUCATION\n"
        "B.S. Computer Science | NYU | 2018\n"
    )
    page.insert_text((72, 72), text, fontsize=11)
    doc.save(pdf_path)
    doc.close()
    yield pdf_path
    if os.path.exists(pdf_path):
        os.remove(pdf_path)


@pytest.fixture
def sample_docx():
    docx_path = tempfile.mktemp(suffix=".docx")
    document = docx.Document()
    document.add_paragraph("John Doe")
    document.add_paragraph("john@example.com | 555-0100 | New York, NY")
    document.add_paragraph("")
    document.add_paragraph("PROFESSIONAL SUMMARY")
    document.add_paragraph("Experienced software engineer with 5+ years in web development.")
    document.add_paragraph("")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, FastAPI, PostgreSQL, Docker")
    document.add_paragraph("")
    document.add_paragraph("EXPERIENCE")
    p = document.add_paragraph()
    p.add_run("Senior Developer").bold = True
    p.add_run(" | Tech Corp | Jan 2020 - Present")
    document.add_paragraph("- Led development of microservices architecture", style="List Bullet")
    document.add_paragraph("- Improved API response time by 40%", style="List Bullet")
    document.add_paragraph("")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("B.S. Computer Science | NYU | 2018")
    document.save(docx_path)
    yield docx_path
    if os.path.exists(docx_path):
        os.remove(docx_path)


@pytest.fixture
def sample_txt():
    txt_path = tempfile.mktemp(suffix=".txt")
    with open(txt_path, "w") as f:
        f.write(
            "John Doe\n"
            "john@example.com | 555-0100\n"
            "\n"
            "SKILLS\n"
            "Python, JavaScript, React\n"
            "\n"
            "EXPERIENCE\n"
            "Frontend Dev | Startup Inc | 2021 - 2023\n"
            "- Built responsive web applications\n"
        )
    yield txt_path
    if os.path.exists(txt_path):
        os.remove(txt_path)


def make_parser():
    with patch("app.services.resume_parser.get_ai_provider", return_value=FakeProvider()):
        return ResumeParserService()


@pytest.mark.asyncio
async def test_extract_pdf_text(sample_pdf):
    parser = make_parser()
    try:
        text = await parser.extract_text(sample_pdf)
        assert "John Doe" in text
        assert "john@example.com" in text
        assert "Python" in text
        assert "Tech Corp" in text
        assert "NYU" in text
    finally:
        await parser.close()


@pytest.mark.asyncio
async def test_extract_docx_text(sample_docx):
    parser = make_parser()
    try:
        text = await parser.extract_text(sample_docx)
        assert "John Doe" in text
        assert "john@example.com" in text
        assert "Python" in text
        assert "Tech Corp" in text
    finally:
        await parser.close()


@pytest.mark.asyncio
async def test_extract_txt_text(sample_txt):
    parser = make_parser()
    try:
        text = await parser.extract_text(sample_txt)
        assert "John Doe" in text
        assert "john@example.com" in text
        assert "Python" in text
    finally:
        await parser.close()


@pytest.mark.asyncio
async def test_extract_unsupported_format():
    parser = make_parser()
    try:
        bad_path = tempfile.mktemp(suffix=".xyz")
        with open(bad_path, "w") as f:
            f.write("test")
        with pytest.raises(ValueError, match="Unsupported file format"):
            await parser.extract_text(bad_path)
        os.remove(bad_path)
    finally:
        await parser.close()


def test_extract_pdf_text_sync(sample_pdf):
    parser = make_parser()
    text = parser._extract_pdf(sample_pdf)
    assert "John Doe" in text
    assert "PROFESSIONAL SUMMARY" in text


def test_extract_docx_text_sync(sample_docx):
    parser = make_parser()
    text = parser._extract_docx(sample_docx)
    assert "John Doe" in text
    assert "PROFESSIONAL SUMMARY" in text


def test_extract_txt_text_sync(sample_txt):
    parser = make_parser()
    text = parser._extract_txt(sample_txt)
    assert "John Doe" in text
    assert "Python" in text
